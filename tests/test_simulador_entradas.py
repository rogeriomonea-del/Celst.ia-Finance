"""A aba Simulador tem entradas próprias — o motor precisa honrá-las.

Na planilha, "Valor inicial" e "Valor mensal" são digitados na própria aba e não
referenciam o Config, então os dois podem divergir. O motor reproduz o que a
planilha mostra e avisa sobre a diferença, em vez de escolher um dos dois em
silêncio.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from engine import extract, pipeline, workbook  # noqa: E402
from engine.models import SimulatorInputs  # noqa: E402


def _planilha(inicial_config: float, inicial_simulador: float, unidade_prazo: str = "ano(s)") -> bytes:
    """Pasta de trabalho sintética com Config + Simulador + uma posição."""
    book = openpyxl.Workbook()

    config = book.active
    config.title = "Config"
    config["A4"], config["B4"] = "Capital inicial (R$)", inicial_config
    config["A5"], config["B5"] = "Aporte mensal (R$)", 1_000.0

    sim = book.create_sheet("Simulador")
    sim["A4"], sim["B4"] = "Valor inicial (R$)", inicial_simulador
    sim["A5"], sim["B5"] = "Valor mensal (R$) [aportes - retiradas]", 1_000.0
    sim["A6"], sim["B6"], sim["C6"] = "Taxa de juros (%)", 12.0, "anual"
    sim["A7"], sim["B7"], sim["C7"] = "Período", 2, unidade_prazo

    posicao = book.create_sheet("Posicao_B3")
    posicao["B8"], posicao["C8"], posicao["D8"] = "Classe", "Ticker", "Ativo"
    posicao["E8"], posicao["F8"] = "Qtde", "Preço (R$)"
    posicao["B9"], posicao["C9"], posicao["D9"] = "Ações", "TEST3", "Teste"
    posicao["E9"], posicao["F9"] = 100, 10.0

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_le_entradas_da_aba_simulador():
    dados = workbook.load(_planilha(500_000.0, 100_000.0), "teste.xlsx")
    entradas = extract.extract_simulator_inputs(dados)

    assert entradas == SimulatorInputs(
        initial=100_000.0, monthly=1_000.0, annual_rate=0.12, months=24
    )


def test_prazo_em_meses_nao_e_multiplicado():
    dados = workbook.load(_planilha(500_000.0, 100_000.0, unidade_prazo="mes(es)"), "teste.xlsx")
    entradas = extract.extract_simulator_inputs(dados)

    assert entradas is not None and entradas.months == 2


def test_simulacao_segue_a_aba_e_avisa_da_divergencia():
    analise = pipeline.analyze_bytes(_planilha(500_000.0, 100_000.0), "teste.xlsx")

    assert analise.simulation is not None
    # Parte dos 100 mil da aba, não dos 500 mil do Config.
    assert analise.simulation.series[0].balance == pytest.approx(100_000.0)
    assert analise.simulation.total_invested == pytest.approx(100_000.0 + 1_000.0 * 24)
    assert any("valor inicial próprio" in aviso for aviso in analise.warnings)


def test_sem_divergencia_nao_gera_aviso():
    analise = pipeline.analyze_bytes(_planilha(100_000.0, 100_000.0), "teste.xlsx")

    assert not any("valor inicial próprio" in aviso for aviso in analise.warnings)


def test_planilha_sem_aba_simulador_usa_o_config():
    book = openpyxl.Workbook()
    config = book.active
    config.title = "Config"
    config["A4"], config["B4"] = "Capital inicial (R$)", 250_000.0
    config["A5"], config["B5"] = "Aporte mensal (R$)", 0.0
    buffer = io.BytesIO()
    book.save(buffer)

    dados = workbook.load(buffer.getvalue(), "teste.xlsx")
    assert extract.extract_simulator_inputs(dados) is None

    analise = pipeline.analyze_bytes(buffer.getvalue(), "teste.xlsx")
    assert analise.simulation is not None
    assert analise.simulation.series[0].balance == pytest.approx(250_000.0)
