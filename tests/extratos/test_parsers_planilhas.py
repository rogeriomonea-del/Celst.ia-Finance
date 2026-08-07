"""Testes dos parsers de planilha (csv_generico, xlsx_generico, xls_legado).

Fixtures 100% sintéticas — nenhum dado real de banco ou cartão.
"""

from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook

from engine.extratos import detect
from engine.extratos.modelos import fingerprint_transacao
from engine.extratos.parsers import csv_generico, xls_legado, xlsx_generico
from engine.extratos.parsers.base import ContextoParse

CSV = csv_generico.ADAPTADORES[0]
XLSX = xlsx_generico.ADAPTADORES[0]
XLS = xls_legado.ADAPTADORES[0]


@pytest.fixture(autouse=True)
def _dados_em_tmp(tmp_path, monkeypatch):
    """Nenhum teste pode encostar no diretório de dados real do usuário."""
    monkeypatch.setenv("CELESTIA_DATA_DIR", str(tmp_path))


def contexto_de(conteudo: bytes, nome: str, **extras) -> ContextoParse:
    return ContextoParse(
        conteudo=conteudo,
        formato=detect.detectar(conteudo, nome),
        nome_original=nome,
        rotulo_origem=extras.pop("rotulo_origem", "conta:Conta Teste"),
        **extras,
    )


# --------------------------------------------------------------------------
# CSV pt-BR
# --------------------------------------------------------------------------

def test_csv_ptbr_ponto_e_virgula_com_virgula_decimal():
    conteudo = (
        "Data;Histórico;Valor;Saldo;Documento\n"
        "05/03/2026;PIX RECEBIDO FULANO;1.234,56;2.234,56;DOC123\n"
        '06/03/2026;"COMPRA\nSUPERMERCADO BOM PREÇO";-200,00;2.034,56;\n'
        "07/03/2026;LOJA MOVEIS PARC 3/10;-150,00;1.884,56;\n"
    ).encode("utf-8")
    contexto = contexto_de(conteudo, "extrato.csv")
    assert contexto.formato.formato == "csv"
    assert CSV.detecta(contexto) > 0

    resultado = CSV.parse(contexto)
    assert resultado.adaptador == "csv-generico"
    assert resultado.mapeamento_necessario is None
    assert resultado.erros == []
    assert len(resultado.transacoes) == 3

    t1, t2, t3 = resultado.transacoes
    assert t1.data_operacao == date(2026, 3, 5)
    assert t1.valor_centavos == 123456
    assert (t1.direcao, t1.tipo) == ("credito", "receita")
    assert t1.saldo_apos_centavos == 223456
    assert t1.documento == "DOC123"
    assert t1.origem_ref == "linha 2"

    # descrição multilinha (célula com \n) vira uma linha só
    assert (t2.direcao, t2.valor_centavos) == ("debito", 20000)
    assert "SUPERMERCADO" in t2.descricao_original
    assert "\n" not in t2.descricao_original

    # a célula multilinha ocupa 2 linhas físicas → t3 começa na linha real 5
    assert t3.origem_ref == "linha 5"
    assert (t3.parcela_num, t3.parcela_total) == (3, 10)
    # a data 07/03/2026 na mesma linha não vira parcela
    assert (t1.parcela_num, t1.parcela_total) == (None, None)

    assert resultado.periodo_inicio == date(2026, 3, 5)
    assert resultado.periodo_fim == date(2026, 3, 7)
    # linhas brutas espelham cabeçalho + 3 registros
    assert len(resultado.linhas_brutas) == 4


def test_csv_internacional_virgula_e_ponto_decimal():
    conteudo = (
        "data,descricao,valor\n"
        '03/04/2026,INTL PAYMENT,"1,234.56"\n'
        "04/04/2026,REFUND,-50.25\n"
    ).encode("utf-8")
    resultado = CSV.parse(contexto_de(conteudo, "statement.csv"))
    assert resultado.erros == []
    assert [t.valor_centavos for t in resultado.transacoes] == [123456, 5025]
    assert [t.direcao for t in resultado.transacoes] == ["credito", "debito"]


def test_csv_cp1252_preserva_acentos():
    texto = "Data;Descrição;Valor\n10/01/2026;CAFÉ SÃO JOÃO;-15,50\n"
    conteudo = texto.encode("cp1252")
    contexto = contexto_de(conteudo, "extrato.csv")
    assert contexto.formato.encoding == "cp1252"
    resultado = CSV.parse(contexto)
    assert resultado.erros == []
    assert resultado.transacoes[0].descricao_original == "CAFÉ SÃO JOÃO"
    assert resultado.transacoes[0].valor_centavos == 1550


def test_csv_colunas_debito_credito_separadas():
    conteudo = (
        "Data;Lançamento;Débito;Crédito\n"
        "02/05/2026;TARIFA MENSAL;25,00;\n"
        "03/05/2026;DEPÓSITO;;300,00\n"
    ).encode("utf-8")
    resultado = CSV.parse(contexto_de(conteudo, "extrato.csv"))
    assert resultado.erros == []
    t_debito, t_credito = resultado.transacoes
    assert (t_debito.direcao, t_debito.valor_centavos, t_debito.tipo) == (
        "debito", 2500, "despesa",
    )
    assert (t_credito.direcao, t_credito.valor_centavos, t_credito.tipo) == (
        "credito", 30000, "receita",
    )


def test_csv_saldo_inicial_antes_do_cabecalho_e_saldo_por_linha():
    conteudo = (
        "Extrato Conta Corrente;;\n"
        "Saldo Anterior;;1.000,00\n"
        "Data;Histórico;Valor;Saldo\n"
        "05/06/2026;PIX ENVIADO;-100,00;900,00\n"
    ).encode("utf-8")
    resultado = CSV.parse(contexto_de(conteudo, "extrato.csv"))
    assert resultado.erros == []
    assert len(resultado.saldos) == 1
    saldo = resultado.saldos[0]
    assert "anterior" in saldo.rotulo
    assert saldo.valor_centavos == 100000
    # sem data na linha do saldo → ancora no início do período
    assert saldo.data == date(2026, 6, 5)
    assert resultado.transacoes[0].saldo_apos_centavos == 90000
    assert resultado.transacoes[0].origem_ref == "linha 4"


def test_csv_transacoes_identicas_ganham_ocorrencias_e_fingerprints_distintas():
    conteudo = (
        "Data;Histórico;Valor\n"
        "05/06/2026;UBER TRIP;-25,00\n"
        "05/06/2026;UBER TRIP;-25,00\n"
    ).encode("utf-8")
    resultado = CSV.parse(contexto_de(conteudo, "extrato.csv"))
    assert len(resultado.transacoes) == 2
    fp1, fp2 = (t.fingerprint for t in resultado.transacoes)
    assert fp1 != fp2
    # e as fingerprints são exatamente as do contrato (ocorrência 1 e 2)
    esperada = lambda n: fingerprint_transacao(  # noqa: E731
        "conta:Conta Teste", date(2026, 6, 5), 2500, "BRL", "UBER TRIP", None, n,
    )
    assert (fp1, fp2) == (esperada(1), esperada(2))


# --------------------------------------------------------------------------
# Mapeamento manual (cabeçalho estranho → 2ª passada)
# --------------------------------------------------------------------------

CSV_CABECALHO_ESTRANHO = (
    "Quando;O que;Quanto\n"
    "01/02/2026;MERCADO;-10,00\n"
    "02/02/2026;SALARIO;5.000,00\n"
).encode("utf-8")


def test_csv_cabecalho_estranho_pede_mapeamento_com_amostra_e_sugestoes():
    resultado = CSV.parse(contexto_de(CSV_CABECALHO_ESTRANHO, "estranho.csv"))
    assert resultado.transacoes == []
    assert resultado.linhas_brutas == []
    pedido = resultado.mapeamento_necessario
    assert pedido is not None
    assert pedido["colunas"] == ["Quando", "O que", "Quanto"]
    assert len(pedido["amostra"]) == 3  # primeiras 8 linhas no máximo
    assert pedido["sugestoes"] == {
        "Quando": "data", "O que": "descricao", "Quanto": "valor",
    }


def test_csv_segunda_passada_com_mapeamento_por_nome():
    contexto = contexto_de(
        CSV_CABECALHO_ESTRANHO,
        "estranho.csv",
        mapeamento={"Quando": "data", "O que": "descricao", "Quanto": "valor"},
    )
    resultado = CSV.parse(contexto)
    assert resultado.mapeamento_necessario is None
    assert resultado.erros == []
    t1, t2 = resultado.transacoes
    assert (t1.data_operacao, t1.direcao, t1.valor_centavos) == (
        date(2026, 2, 1), "debito", 1000,
    )
    assert (t2.direcao, t2.valor_centavos) == ("credito", 500000)


def test_csv_segunda_passada_com_mapeamento_por_indice():
    contexto = contexto_de(
        CSV_CABECALHO_ESTRANHO,
        "estranho.csv",
        mapeamento={"0": "data", "1": "descricao", "2": "valor"},
    )
    resultado = CSV.parse(contexto)
    assert len(resultado.transacoes) == 2
    assert resultado.transacoes[0].descricao_original == "MERCADO"


def test_csv_mapeamento_incompleto_volta_a_pedir_mapeamento():
    contexto = contexto_de(
        CSV_CABECALHO_ESTRANHO, "estranho.csv", mapeamento={"1": "descricao"},
    )
    resultado = CSV.parse(contexto)
    assert resultado.transacoes == []
    assert resultado.mapeamento_necessario is not None
    assert any("mapeamento incompleto" in aviso for aviso in resultado.avisos)


# --------------------------------------------------------------------------
# Neutralização de injeção de fórmula
# --------------------------------------------------------------------------

def test_csv_formula_maliciosa_neutralizada_no_espelho():
    conteudo = (
        "Data;Histórico;Valor\n"
        "01/08/2026;=cmd()|whoami;-10,00\n"
    ).encode("utf-8")
    resultado = CSV.parse(contexto_de(conteudo, "extrato.csv"))
    celulas = resultado.linhas_brutas[1].conteudo["celulas"]
    assert celulas[1] == "'=cmd()|whoami"
    # número negativo legítimo NÃO ganha apóstrofo
    assert celulas[2] == "-10,00"
    # e o valor monetário parseia normalmente
    assert resultado.transacoes[0].valor_centavos == 1000


# --------------------------------------------------------------------------
# XLSX multiabas
# --------------------------------------------------------------------------

def _xlsx_sintetico() -> bytes:
    pasta = Workbook()
    resumo = pasta.active
    resumo.title = "Resumo"
    resumo.append(["Relatório do período"])
    resumo.append(["Total", 123])

    conta = pasta.create_sheet("Conta")
    conta.append(["Data", "Histórico", "Valor"])
    conta.append([datetime(2026, 7, 1), "PIX RECEBIDO", 1500.5])
    conta.append(["02/07/2026", "TARIFA", "-30,00"])

    oculta = pasta.create_sheet("Oculta")
    oculta.append(["Data", "Histórico", "Valor"])
    oculta.append(["01/01/2020", "FANTASMA", "-1,00"])
    oculta.sheet_state = "hidden"

    buffer = io.BytesIO()
    pasta.save(buffer)
    return buffer.getvalue()


def test_xlsx_multiabas_processa_visiveis_e_ignora_ocultas():
    conteudo = _xlsx_sintetico()
    contexto = contexto_de(conteudo, "extrato.xlsx")
    assert contexto.formato.formato == "xlsx"
    assert XLSX.detecta(contexto) > 0

    resultado = XLSX.parse(contexto)
    assert resultado.adaptador == "xlsx-generico"
    assert resultado.mapeamento_necessario is None
    assert len(resultado.transacoes) == 2

    t1, t2 = resultado.transacoes
    assert t1.origem_ref == "Conta!L2"
    assert t1.data_operacao == date(2026, 7, 1)
    assert (t1.direcao, t1.valor_centavos) == ("credito", 150050)
    assert t2.origem_ref == "Conta!L3"
    assert (t2.direcao, t2.valor_centavos) == ("debito", 3000)

    # a aba oculta nunca entra; a aba sem cabeçalho vira aviso
    assert all("FANTASMA" not in t.descricao_original for t in resultado.transacoes)
    assert any("Resumo" in aviso for aviso in resultado.avisos)


def test_xlsx_sem_nenhum_cabecalho_pede_mapeamento():
    pasta = Workbook()
    aba = pasta.active
    aba.title = "Dados"
    aba.append(["Quando", "Quanto"])
    aba.append(["01/02/2026", "-10,00"])
    buffer = io.BytesIO()
    pasta.save(buffer)
    resultado = XLSX.parse(contexto_de(buffer.getvalue(), "estranho.xlsx"))
    assert resultado.transacoes == []
    assert resultado.mapeamento_necessario is not None
    assert resultado.mapeamento_necessario["aba"] == "Dados"


def test_xlsx_corrompido_vira_erro_e_nao_excecao():
    conteudo = b"PK\x03\x04lixo-que-nao-e-zip-valido"
    resultado = XLSX.parse(
        ContextoParse(
            conteudo=conteudo,
            formato=detect.FormatoDetectado("xlsx", "application/zip"),
            nome_original="quebrado.xlsx",
        )
    )
    assert resultado.transacoes == []
    assert resultado.erros


# --------------------------------------------------------------------------
# XLS legado: HTML disfarçado e BIFF binário
# --------------------------------------------------------------------------

def test_xls_html_de_banco_com_entidades():
    conteudo = (
        "<html><body>"
        "<table><tr><td>Banco Sint&eacute;tico S.A.</td></tr></table>"
        "<table>"
        "<tr><th>Data</th><th>Hist&oacute;rico</th><th>Valor</th></tr>"
        "<tr><td>10/07/2026</td><td>COMPRA C&amp;A</td><td>-99,90</td></tr>"
        "<tr><td>11/07/2026</td><td>PIX RECEBIDO<br>JO&Atilde;O</td>"
        "<td>150,00</td></tr>"
        "</table>"
        "</body></html>"
    ).encode("utf-8")
    contexto = contexto_de(conteudo, "extrato.xls")
    assert contexto.formato.formato == "xls_html"
    assert XLS.detecta(contexto) > 0

    resultado = XLS.parse(contexto)
    assert resultado.adaptador == "xls-legado"
    assert resultado.erros == []
    assert len(resultado.transacoes) == 2
    t1, t2 = resultado.transacoes
    assert t1.descricao_original == "COMPRA C&A"  # &amp; decodificado
    assert (t1.direcao, t1.valor_centavos) == ("debito", 9990)
    assert "JOÃO" in t2.descricao_original  # &Atilde; decodificado, <br> unido
    assert (t2.direcao, t2.valor_centavos) == ("credito", 15000)


def test_xls_biff_binario_recusado_com_aviso_claro():
    conteudo = bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 64
    contexto = contexto_de(conteudo, "extrato.xls")
    assert contexto.formato.formato == "xls_biff"

    resultado = XLS.parse(contexto)
    assert resultado.transacoes == []
    assert xls_legado.AVISO_BIFF in resultado.avisos
    assert "exporte como XLSX ou CSV" in resultado.avisos[0]
    assert resultado.erros and "XLSX ou CSV" in resultado.erros[0]["mensagem"]
