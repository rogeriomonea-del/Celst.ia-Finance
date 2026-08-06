"""Leitura de pastas de trabalho (.xlsm/.xlsx/.csv) sem depender do Excel.

O motor lê os *valores* das células — nunca reavalia fórmulas do arquivo.
Toda a lógica que no Excel morava em fórmulas foi reescrita em ``engine.calc``.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Dict, Iterable, List, Optional, Tuple, Union

Source = Union[str, bytes, io.BytesIO]

_A1 = re.compile(r"^([A-Za-z]+)(\d+)$")


def column_index(letters: str) -> int:
    """'A' -> 1, 'B' -> 2, 'AA' -> 27."""
    idx = 0
    for ch in letters.upper():
        idx = idx * 26 + (ord(ch) - 64)
    return idx


def column_letter(index: int) -> str:
    letters = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


class SheetGrid:
    """Acesso indexado às células de uma aba, tolerante a buracos."""

    __slots__ = ("title", "_cells", "max_row", "max_col")

    def __init__(self, title: str, cells: Dict[Tuple[int, int], Any]):
        self.title = title
        self._cells = cells
        self.max_row = max((r for r, _ in cells), default=0)
        self.max_col = max((c for _, c in cells), default=0)

    def value(self, row: int, col: int) -> Any:
        return self._cells.get((row, col))

    def cell(self, ref: str) -> Any:
        m = _A1.match(ref.replace("$", ""))
        if not m:
            raise ValueError(f"referência inválida: {ref}")
        return self.value(int(m.group(2)), column_index(m.group(1)))

    def text(self, row: int, col: int) -> str:
        value = self.value(row, col)
        if value is None:
            return ""
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return str(value).strip()

    def number(self, row: int, col: int, default: float = 0.0) -> float:
        return as_number(self.value(row, col), default)

    def date(self, row: int, col: int) -> Optional[date]:
        return as_date(self.value(row, col))

    def row_values(self, row: int) -> List[Any]:
        return [self.value(row, c) for c in range(1, self.max_col + 1)]

    def iter_rows(self, start: int = 1, end: Optional[int] = None) -> Iterable[int]:
        return range(start, (end or self.max_row) + 1)

    def find_row(self, needle: str, col: Optional[int] = None, limit: int = 400) -> Optional[int]:
        """Primeira linha cujo texto (em ``col`` ou em qualquer coluna) casa."""
        target = needle.strip().lower()
        cols = [col] if col else range(1, min(self.max_col, 30) + 1)
        for row in range(1, min(self.max_row, limit) + 1):
            for c in cols:
                if self.text(row, c).strip().lower() == target:
                    return row
        return None

    def header_map(self, row: int) -> Dict[str, int]:
        """{cabeçalho normalizado: índice da coluna} para a linha informada."""
        headers: Dict[str, int] = {}
        for col in range(1, self.max_col + 1):
            label = normalize(self.text(row, col))
            if label and label not in headers:
                headers[label] = col
        return headers


@dataclass
class WorkbookData:
    """Todas as abas de um arquivo, com os valores já materializados."""

    file_name: str
    sheets: Dict[str, SheetGrid] = field(default_factory=dict)
    kind: str = "xlsx"
    warnings: List[str] = field(default_factory=list)

    def sheet(self, *names: str) -> Optional[SheetGrid]:
        """Primeira aba que casar (comparação normalizada e tolerante)."""
        wanted = [normalize(n) for n in names]
        table = {normalize(title): grid for title, grid in self.sheets.items()}
        for name in wanted:
            if name in table:
                return table[name]
        for name in wanted:
            for title, grid in table.items():
                if name and (name in title or title in name):
                    return grid
        return None

    @property
    def titles(self) -> List[str]:
        return list(self.sheets)


# --------------------------------------------------------------------------
# Conversores
# --------------------------------------------------------------------------

_NUMBER_CLEAN = re.compile(r"[^\d,.\-]")


def as_number(value: Any, default: float = 0.0) -> float:
    """Converte para float aceitando formatos pt-BR ('1.234,56', 'R$ 8,5%')."""
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return default
    percent = text.endswith("%")
    text = _NUMBER_CLEAN.sub("", text)
    if not text or text in ("-", ".", ","):
        return default
    if "," in text and "." in text:
        # 1.234,56 (pt-BR) vs 1,234.56 (en-US): separador decimal é o último.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return default
    return number / 100 if percent else number


def as_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


_ACCENTS = str.maketrans("áàâãäéèêëíìîïóòôõöúùûüçñ", "aaaaaeeeeiiiiooooouuuucn")


def normalize(text: str) -> str:
    """Minúsculas, sem acento, sem pontuação — para casar cabeçalhos."""
    if not text:
        return ""
    cleaned = str(text).strip().lower().translate(_ACCENTS)
    cleaned = re.sub(r"[^a-z0-9%]+", " ", cleaned)
    return re.sub(r"\s+", " ", cleaned).strip()


# --------------------------------------------------------------------------
# Carregamento
# --------------------------------------------------------------------------


def load(source: Source, file_name: str = "planilha.xlsx") -> WorkbookData:
    """Lê .xlsm/.xlsx (openpyxl) ou .csv/.tsv (stdlib) para ``WorkbookData``."""
    lower = file_name.lower()
    if lower.endswith((".csv", ".tsv", ".txt")):
        return _load_csv(source, file_name)
    return _load_excel(source, file_name)


def _as_stream(source: Source) -> Union[str, io.BytesIO]:
    if isinstance(source, bytes):
        return io.BytesIO(source)
    return source


def _load_excel(source: Source, file_name: str) -> WorkbookData:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - ambiente sem dependência
        raise RuntimeError(
            "openpyxl é necessário para ler .xlsx/.xlsm (pip install openpyxl)"
        ) from exc

    import warnings as _warnings

    with _warnings.catch_warnings():
        _warnings.simplefilter("ignore")
        book = openpyxl.load_workbook(
            _as_stream(source), data_only=True, read_only=True, keep_links=False
        )
    data = WorkbookData(file_name=file_name, kind="xlsm" if file_name.lower().endswith(".xlsm") else "xlsx")
    try:
        for ws in book.worksheets:
            cells: Dict[Tuple[int, int], Any] = {}
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cells[(cell.row, cell.column)] = cell.value
            data.sheets[ws.title] = SheetGrid(ws.title, cells)
    finally:
        book.close()
    return data


def _load_csv(source: Source, file_name: str) -> WorkbookData:
    if isinstance(source, bytes):
        raw = source
    elif isinstance(source, io.BytesIO):
        raw = source.getvalue()
    else:
        with open(source, "rb") as handle:
            raw = handle.read()

    text = ""
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","

    cells: Dict[Tuple[int, int], Any] = {}
    for r, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        for c, raw_value in enumerate(row, start=1):
            value = raw_value.strip()
            if value:
                cells[(r, c)] = value
    title = file_name.rsplit("/", 1)[-1].rsplit(".", 1)[0] or "Dados"
    return WorkbookData(file_name=file_name, kind="csv", sheets={title: SheetGrid(title, cells)})
